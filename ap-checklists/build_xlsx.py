# -*- coding: utf-8 -*-
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = "/home/user/Games/ap-checklists/AP-Run-Cards.xlsx"

# ---- palette (matches the web version) ----
PETROL   = "0E4F52"
PETROL_L = "DCEAE9"
OCHRE    = "8F5405"
OCHRE_L  = "F6EBDA"
OXBLOOD  = "8E2226"
OXBLOOD_L= "F7E5E4"
GREEN    = "2C6349"
GREEN_L  = "DFEDE5"
INK      = "141D1B"
MUTED    = "6B7B77"
LINE     = "D3DBD8"
INPUT_F  = "FFF7DC"   # cells you fill in
WHITE    = "FFFFFF"

F = "Arial"
def font(sz=10, b=False, color=INK, i=False): return Font(name=F, size=sz, bold=b, color=color, italic=i)
def fill(c): return PatternFill("solid", fgColor=c)
thin = Side(style="thin", color=LINE)
box  = Border(left=thin, right=thin, top=thin, bottom=thin)
bot  = Border(bottom=thin)

WRAP   = Alignment(wrap_text=True, vertical="top")
WRAPC  = Alignment(wrap_text=True, vertical="center")
CTR    = Alignment(horizontal="center", vertical="center")

TICKS = '"✓,—"'   # ✓ = done, — = not applicable

# ---------------- content ----------------
CARDS = [
 dict(id="A", tab="A Review", title="Invoice review gate", style="Do-confirm", cadence="Per invoice",
   howto=("Review the invoice your normal way. Then stop at the same moment every time and confirm "
          "these. Because this card runs dozens of times a day it is deliberately short — resist "
          "adding to it."),
   pause="Your pause point:", pause_val="[e.g. before you change status to Approved for Payment]",
   steps=[
    ("Duplicate check. Same vendor + same invoice number. Then also same vendor + same amount + near date — duplicates slip through when the invoice number is keyed differently.", "[Where you check]", 1),
    ("Approval is present, and the approver has authority for this amount.", "[Your approval limits]", 1),
    ("Not on hold, and not in dispute.", "[Where you check]", 0),
    ("Open credit memos for this vendor checked, and applied or netted.", "", 0),
    ("Terms and due date match the vendor master, not just what the invoice says. Note any early-pay discount and the date it expires.", "", 0),
    ("Remit-to on the invoice matches the vendor master. Any difference at all: stop and run Card E. Do not fix it inline.", "Difference → Card E", 1),
    ("Correct entity or company code.", "[If you have more than one]", 0),
    ("Set status to ready-to-schedule.", "[Your status name]", 0),
    ("[Add your own step — especially anything upstream gets wrong often enough that you catch it]", "", 0),
   ]),
 dict(id="B", tab="B Run build", title="Run build", style="Read-do", cadence="Per run",
   howto="Read each step, then do it. This card ends by recording the batch total and item count that Card C reconciles against.",
   pause="Run cadence:", pause_val="[e.g. every Tuesday and Friday]",
   inputs=[("Run date", "b_date", "[dd/mm/yyyy]"), ("Batch total", "b_total", ""), ("Item count", "b_count", "")],
   inputs_title="RECORD BEFORE YOU RELEASE — Card C checks against these",
   steps=[
    ("Pull the payment proposal.", "[Your date range]", 0),
    ("Pull-forward check: anything due before the NEXT run after this one has to go in this run. Compare due dates to the next run date, not to today.", "[Next run date]", 0),
    ("Discount check: any early-pay discount expiring before the next run. A lost discount is a silent error — nobody reports it.", "", 0),
    ("Exclusions: remove anything on hold, in dispute, pending vendor verification, or missing approval.", "", 0),
    ("Credit memos applied, so vendors net correctly instead of paying gross.", "", 0),
    ("Payment method per vendor is correct. Watch for any vendor with more than one active method — that is how one invoice goes out twice by two channels.", "[Methods you use]", 0),
    ("Funding or cash check.", "[Confirm funds, or hand a total to Treasury?]", 0),
    ("Batch total and item count recorded in the box above, and the run parked as built-not-released.", "[Your status name]", 1),
   ]),
 dict(id="C", tab="C Release", title="Release", style="Read-do", cadence="Per run",
   howto="The shortest and strictest card. Do not do this one from memory, and do not do it while doing anything else.",
   stop_top=("THE RE-TRANSMIT RULE",
     "If you were interrupted and you are not certain the file went out, DO NOT RE-SEND. Open the bank portal and "
     "look at today's activity. The bank's record is the only source of truth — your memory of the last ten minutes "
     "is not, and re-sending duplicates every payment in the run."),
   recon=True,
   steps=[
    ("Batch total on screen matches the figure recorded in Card B. If it does not match, stop and find out why — never adjust the recorded figure to match the screen.", "", 1),
    ("Item count on screen matches the figure recorded in Card B.", "", 1),
    ("Funding account is correct.", "[Correct account]", 1),
    ("Effective or value date is correct.", "", 0),
    ("Second review and release.", "[Who? If nobody, see note below]", 0),
    ("Transmit.", "", 0),
    ("Bank confirmation captured — reference, timestamp, screenshot — and recorded below.", "[Where you save it]", 1),
    ("Positive pay file uploaded.", "[If applicable]", 0),
    ("Check stock sequence recorded.", "[If applicable]", 0),
   ],
   inputs=[("Confirmation ref", "c_ref", ""), ("Time sent", "c_time", ""), ("Check stock from / to", "c_stock", "")],
   inputs_title="RECORD AT RELEASE",
   note=("IF YOU ARE IN SOLE CONTROL",
     "Building and releasing runs with nobody else in the loop is a standard segregation-of-duties gap. It is worth "
     "raising with your manager, and it is not yours to solve alone. Until then, put real time between the two steps: "
     "build the run, do something else, come back to release it. The gap is a partial substitute for a second pair of eyes.")),
 dict(id="D", tab="D Post-run", title="Post-run", style="Read-do", cadence="Same day as the run",
   howto="The quiet failures nobody notices for a month. Nothing here can double-send money, so it is safe to be interrupted — just finish it before end of day.",
   steps=[
    ("Confirmation filed.", "[Where]", 0),
    ("Payment statuses updated. Nothing left sitting in released limbo.", "[Your system]", 0),
    ("Remittance advice sent to vendors.", "[Automatic or manual — how?]", 0),
    ("Run reconciled to the bank.", "[Same day or next day?]", 0),
    ("Rejects and returns worked the same day — not when you get to it. A returned ACH that sits is a vendor who thinks they were paid and wasn't.", "[Your returns procedure]", 0),
    ("Re-check the hold pile: anything excluded in Card B is still correctly parked and hasn't quietly aged past due.", "[How you re-check it]", 0),
    ("Run logged on the Run log tab.", "", 0),
    ("[Add anything your close needs from you at run time rather than at month end]", "", 0),
   ]),
 dict(id="E", tab="E Bank change", title="Vendor bank or remit-to change", style="Read-do", cadence="Triggered, never scheduled",
   howto="Trigger: any new vendor, or any change to remit-to address, bank account or routing number — however it arrives, and however routine it looks.",
   stop_top=("WHY THIS CARD IS RIGID",
     "This is the highest-consequence, lowest-frequency thing you touch. Business email compromise enters accounts "
     "payable almost entirely at this step, and it works precisely because the request arrives when you are busy and "
     "looks completely normal. Urgency in the request is a warning sign, not a reason to hurry."),
   steps=[
    ("Hold all payments to this vendor until this card is finished.", "", 1),
    ("Verify by phone, using a number already on file — in the vendor master or on a prior invoice. NEVER a number from the email, attachment, or signature block requesting the change.", "", 1),
    ("Speak to a known contact, or someone whose identity you can confirm independently.", "", 1),
    ("Read the new digits back to them. Don't ask \"did you send a change?\" — ask them to state the account.", "", 1),
    ("Second person confirms the change before it is entered.", "[Who]", 1),
    ("Verification documented below.", "", 0),
    ("Change entered.", "[Your system]", 0),
    ("Vendor's next payment held one cycle, then released.", "[Or your policy]", 1),
    ("Documentation filed.", "[Where]", 0),
   ],
   inputs=[("Vendor", "e_vendor", ""), ("Number called (and where it came from)", "e_num", ""),
           ("Spoke with", "e_who", ""), ("Date and time of call", "e_when", ""), ("Second reviewer", "e_second", "")],
   inputs_title="VERIFICATION RECORD",
   note=("ESCALATE IMMEDIATELY IF",
     "• The number on file doesn't work, or goes somewhere unexpected   • The contact is new, or unknown to you   "
     "• The request is urgent, or pushes back on being verified   • The email domain differs from the vendor's usual "
     "one, even by a character   • The change arrives right before a large scheduled payment\n"
     "Escalate to [who]. Do not process. It is always acceptable to be slow here.")),
]

NEVER_SKIP = [
 "Duplicate check",
 "Approval present, and from someone with authority for the amount",
 "Remit-to matches the vendor master",
 "Batch total and count reconciled at release",
 "Funding account correct",
]

TRIGGERS = [
 ("New vendor, or a changed remit-to or bank detail", "Card E"),
 ("Invoice is on hold or in dispute", "[your hold procedure]"),
 ("Open credit memo for this vendor", "[your credit memo procedure]"),
 ("Amount over [$ threshold]", "[extra approval]"),
 ("Payment returned or rejected by the bank", "[your returns procedure]"),
 ("Vendor asks for an urgent, off-cycle payment", "[your off-cycle procedure]"),
]

wb = Workbook()

# ============ card sheets ============
def rowh(text, chars_per_line=78, base=14.5, pad=6):
    lines = max(1, math.ceil(len(text) / chars_per_line))
    return max(20, lines * base + pad)

meta = {}

def build_card(c):
    ws = wb.create_sheet(c["tab"])
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 5.5
    ws.column_dimensions["B"].width = 74
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 13
    r = 1

    # header band
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    hc = ws.cell(r, 1, "CARD " + c["id"] + "   ·   " + c["title"].upper())
    hc.font = font(13, True, WHITE); hc.fill = fill(PETROL); hc.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r].height = 32
    for col in range(1, 5): ws.cell(r, col).fill = fill(PETROL)
    r += 1

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    sc = ws.cell(r, 1, c["style"].upper() + "   ·   " + c["cadence"].upper())
    sc.font = font(8, True, PETROL); sc.fill = fill(PETROL_L); sc.alignment = Alignment(vertical="center", indent=1)
    for col in range(1, 5): ws.cell(r, col).fill = fill(PETROL_L)
    ws.row_dimensions[r].height = 18
    r += 2

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    hw = ws.cell(r, 1, c["howto"]); hw.font = font(10, color=MUTED); hw.alignment = WRAP
    ws.row_dimensions[r].height = rowh(c["howto"], 150)
    r += 1

    if c.get("pause"):
        ws.cell(r, 1, "").font = font()
        pl = ws.cell(r, 2, c["pause"]); pl.font = font(10, True); pl.alignment = Alignment(vertical="center")
        pv = ws.cell(r, 3, c["pause_val"]); pv.font = font(10); pv.fill = fill(INPUT_F)
        pv.border = box; pv.alignment = WRAPC
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1

    # stop banner
    if c.get("stop_top"):
        h, body = c["stop_top"]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        t = ws.cell(r, 1, h); t.font = font(10, True, OXBLOOD); t.fill = fill(OXBLOOD_L)
        t.alignment = Alignment(vertical="center", indent=1); t.border = box
        for col in range(1, 5): ws.cell(r, col).fill = fill(OXBLOOD_L); ws.cell(r, col).border = box
        ws.row_dimensions[r].height = 20; r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        b = ws.cell(r, 1, body); b.font = font(10, color=OXBLOOD); b.fill = fill(OXBLOOD_L)
        b.alignment = Alignment(wrap_text=True, vertical="top", indent=1); b.border = box
        for col in range(1, 5): ws.cell(r, col).fill = fill(OXBLOOD_L); ws.cell(r, col).border = box
        ws.row_dimensions[r].height = rowh(body, 140, pad=10); r += 2

    # Card B input box (before steps)
    if c["id"] == "B":
        r = input_box(ws, r, c["inputs_title"], c["inputs"], meta_prefix="B") + 1

    # Card C reconciliation box
    if c.get("recon"):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        t = ws.cell(r, 1, "RECONCILE TO CARD B"); t.font = font(9, True, PETROL); t.fill = fill(PETROL_L)
        t.alignment = Alignment(vertical="center", indent=1); t.border = box
        for col in range(1, 5): ws.cell(r, col).fill = fill(PETROL_L); ws.cell(r, col).border = box
        ws.row_dimensions[r].height = 18; r += 1
        hdr = ["", "", "Recorded in Card B", "On screen now"]
        for i, hh in enumerate(hdr, start=1):
            cc = ws.cell(r, i, hh); cc.font = font(9, True, MUTED); cc.alignment = CTR; cc.border = bot
        ws.row_dimensions[r].height = 18; r += 1
        for label, bref in (("Batch total", meta["B_b_total"]), ("Item count", meta["B_b_count"])):
            ws.cell(r, 2, label).font = font(10, True)
            ws.cell(r, 2).alignment = Alignment(vertical="center", indent=1)
            fr = ws.cell(r, 3, "='B Run build'!" + bref); fr.font = font(10, color=GREEN)
            fr.alignment = CTR; fr.border = box
            en = ws.cell(r, 4, None); en.fill = fill(INPUT_F); en.border = box; en.alignment = CTR; en.font = font(10)
            meta["C_screen_" + label.split()[1]] = get_column_letter(4) + str(r)
            meta["C_bref_" + label.split()[1]] = "'B Run build'!" + bref
            ws.row_dimensions[r].height = 24; r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        chk = ws.cell(r, 1,
          '=IF(AND(COUNT({t1},{t2})=2,COUNT({c1},{c2})=2),'
          'IF(AND({t1}={c1},{t2}={c2}),"MATCH — cleared to transmit",'
          '"DOES NOT MATCH — STOP. Find out why. Never adjust the recorded figure to fit the screen."),'
          '"Fill in both columns above before you transmit.")'.format(
            t1=meta["C_bref_total"], t2=meta["C_bref_count"],
            c1=meta["C_screen_total"], c2=meta["C_screen_count"]))
        chk.font = font(11, True); chk.alignment = Alignment(vertical="center", horizontal="center")
        chk.border = box
        rng = "A{0}:D{0}".format(r)
        ws.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("MATCH — cleared",$A{0}))'.format(r)],
            fill=fill(GREEN_L), font=Font(name=F, size=11, bold=True, color=GREEN), stopIfTrue=False))
        ws.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("DOES NOT MATCH",$A{0}))'.format(r)],
            fill=fill(OXBLOOD_L), font=Font(name=F, size=11, bold=True, color=OXBLOOD), stopIfTrue=False))
        ws.row_dimensions[r].height = 30
        r += 2

    # steps header
    prog_row = r
    ws.cell(r, 1, "✓").font = font(9, True, MUTED); ws.cell(r, 1).alignment = CTR
    ws.cell(r, 2, "STEP").font = font(9, True, MUTED)
    ws.cell(r, 3, "YOUR DETAIL — edit the yellow cells").font = font(9, True, MUTED)
    ws.cell(r, 4, "").font = font(9)
    for col in range(1, 5): ws.cell(r, col).border = bot
    ws.row_dimensions[r].height = 20
    r += 1

    first = r
    for text, detail, never in c["steps"]:
        tk = ws.cell(r, 1, None); tk.alignment = CTR; tk.border = box; tk.font = font(12, True, GREEN)
        st = ws.cell(r, 2, text); st.font = font(10); st.alignment = WRAP; st.border = bot
        dt = ws.cell(r, 3, detail if detail else None); dt.alignment = WRAPC; dt.border = bot
        dt.font = font(9, color=MUTED if detail.startswith("[") else INK)
        if detail.startswith("["): dt.fill = fill(INPUT_F)
        ns = ws.cell(r, 4, "NEVER SKIP" if never else None)
        ns.font = font(8, True, OCHRE); ns.alignment = CTR; ns.border = bot
        if never: ns.fill = fill(OCHRE_L)
        ws.row_dimensions[r].height = rowh(text, 82)
        r += 1
    last = r - 1

    dv = DataValidation(type="list", formula1=TICKS, allow_blank=True, showDropDown=False)
    dv.error = "Pick ✓ when the step is done, or — if it does not apply."
    dv.prompt = "✓ = done   — = not applicable"
    ws.add_data_validation(dv)
    dv.add("A{}:A{}".format(first, last))

    # grey out completed step text
    ws.conditional_formatting.add("B{}:D{}".format(first, last),
        FormulaRule(formula=['$A{}="✓"'.format(first)], font=Font(name=F, size=10, color=MUTED), stopIfTrue=False))

    rng = "A{}:A{}".format(first, last)
    meta[c["id"] + "_range"] = "'" + c["tab"] + "'!" + rng.replace("A", "$A")
    meta[c["id"] + "_n"] = len(c["steps"])

    # progress readout in the header row
    p = ws.cell(prog_row, 4, '=COUNTIF({r},"✓")&" of "&({n}-COUNTIF({r},"—"))'.format(r=rng, n=len(c["steps"])))
    p.font = font(10, True, PETROL); p.alignment = CTR
    r += 1

    # trailing input box
    if c.get("inputs") and c["id"] != "B":
        r = input_box(ws, r, c["inputs_title"], c["inputs"], meta_prefix=c["id"]) + 1

    # trailing note
    if c.get("note"):
        h, body = c["note"]
        col = OXBLOOD if c["id"] == "E" else MUTED
        bg  = OXBLOOD_L if c["id"] == "E" else WHITE
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        t = ws.cell(r, 1, h); t.font = font(9, True, col); t.alignment = Alignment(vertical="center", indent=1)
        t.fill = fill(bg); t.border = box
        for cc2 in range(1, 5): ws.cell(r, cc2).fill = fill(bg); ws.cell(r, cc2).border = box
        ws.row_dimensions[r].height = 18; r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        b = ws.cell(r, 1, body); b.font = font(9.5, color=col if c["id"] == "E" else INK)
        b.alignment = Alignment(wrap_text=True, vertical="top", indent=1); b.fill = fill(bg); b.border = box
        for cc2 in range(1, 5): ws.cell(r, cc2).fill = fill(bg); ws.cell(r, cc2).border = box
        ws.row_dimensions[r].height = rowh(body, 130, pad=12); r += 1

    ws.freeze_panes = "A" + str(first)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws

def input_box(ws, r, title, inputs, meta_prefix):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    t = ws.cell(r, 1, title); t.font = font(9, True, PETROL); t.fill = fill(PETROL_L)
    t.alignment = Alignment(vertical="center", indent=1); t.border = box
    for col in range(1, 5): ws.cell(r, col).fill = fill(PETROL_L); ws.cell(r, col).border = box
    ws.row_dimensions[r].height = 18
    r += 1
    for label, key, hint in inputs:
        lb = ws.cell(r, 2, label); lb.font = font(10, True); lb.alignment = Alignment(vertical="center", indent=1)
        lb.border = bot
        vc = ws.cell(r, 3, hint if hint else None)
        vc.fill = fill(INPUT_F); vc.border = box; vc.alignment = WRAPC
        vc.font = font(10, color=MUTED if hint else INK)
        ws.cell(r, 1).border = bot; ws.cell(r, 4).border = bot
        meta[meta_prefix + "_" + key] = "$C$" + str(r)
        ws.row_dimensions[r].height = 26
        r += 1
    return r

for c in CARDS:
    build_card(c)

# ============ Start sheet ============
ws = wb["Sheet"]; ws.title = "Start"
wb.move_sheet("Start", offset=-len(CARDS))
ws.sheet_view.showGridLines = False
for col, w in (("A", 3), ("B", 46), ("C", 30), ("D", 22), ("E", 14)):
    ws.column_dimensions[col].width = w

r = 1
ws.merge_cells("A1:E1")
h = ws.cell(1, 1, "AP RUN CARDS"); h.font = font(16, True, WHITE); h.fill = fill(PETROL)
h.alignment = Alignment(vertical="center", indent=1)
for col in range(1, 6): ws.cell(1, col).fill = fill(PETROL)
ws.row_dimensions[1].height = 38
ws.merge_cells("A2:E2")
s = ws.cell(2, 1, "REVIEW  ·  SCHEDULE  ·  PAY"); s.font = font(8, True, PETROL); s.fill = fill(PETROL_L)
s.alignment = Alignment(vertical="center", indent=1)
for col in range(1, 6): ws.cell(2, col).fill = fill(PETROL_L)
ws.row_dimensions[2].height = 18
r = 4

def heading(ws, r, text, span=5):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
    c = ws.cell(r, 1, text.upper()); c.font = font(9, True, MUTED)
    c.alignment = Alignment(vertical="bottom"); c.border = bot
    for col in range(1, span + 1): ws.cell(r, col).border = bot
    ws.row_dimensions[r].height = 22
    return r + 1

ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
lg = ws.cell(r, 1, "Run one card at a time, never two. Tick column A with the ✓ dropdown. "
                   "Yellow cells are yours to fill in — replace the [bracketed prompts] with how it actually "
                   "works at your shop, once, and they stay. Clear the ✓ column to start a new run.")
lg.font = font(10, color=INK); lg.alignment = WRAP
ws.row_dimensions[r].height = 46
r += 2

r = heading(ws, r, "The run, in order")
for i, cd in enumerate(CARDS):
    letter = ws.cell(r, 1, cd["id"]); letter.font = font(11, True, PETROL); letter.alignment = CTR
    nm = ws.cell(r, 2, cd["title"]); nm.font = font(11, True)
    nm.alignment = Alignment(vertical="center", indent=1)
    st = ws.cell(r, 3, cd["style"] + "  ·  " + cd["cadence"]); st.font = font(9, color=MUTED)
    st.alignment = Alignment(vertical="center")
    pr = ws.cell(r, 4, '=COUNTIF({r},"✓")&" of "&({n}-COUNTIF({r},"—"))&" done"'.format(
        r=meta[cd["id"] + "_range"], n=meta[cd["id"] + "_n"]))
    pr.font = font(10, color=GREEN); pr.alignment = Alignment(vertical="center", horizontal="right")
    lk = ws.cell(r, 5, "Open →"); lk.font = font(10, color=PETROL, b=True)
    lk.hyperlink = "#'" + cd["tab"] + "'!A1"; lk.alignment = CTR
    for col in range(1, 6): ws.cell(r, col).border = bot
    ws.row_dimensions[r].height = 24
    if cd["id"] == "D":
        r += 1
        sep = ws.cell(r, 2, "Card E is triggered, not scheduled — it runs when something sets it off, never in sequence.")
        sep.font = font(9, True, OCHRE, i=True); sep.alignment = Alignment(vertical="center", indent=1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 22
    r += 1
r += 1

r = heading(ws, r, "Never skip these — even on the worst day")
for t in NEVER_SKIP:
    m = ws.cell(r, 1, "■"); m.font = font(9, True, OCHRE); m.alignment = CTR
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    c = ws.cell(r, 2, t); c.font = font(10.5); c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r].height = 21
    r += 1
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
n = ws.cell(r, 2, "Everything else on every card is hygiene and can flex under pressure. These five don't. "
                  "Replace this list once you know what has actually gone wrong for you — it should be evidence, not a guess.")
n.font = font(9, color=MUTED); n.alignment = WRAP
ws.row_dimensions[r].height = 30
r += 2

r = heading(ws, r, "If you see this, stop and run that")
for cond, act in TRIGGERS:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c1 = ws.cell(r, 1, cond); c1.font = font(10); c1.alignment = Alignment(vertical="center", indent=1)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    c2 = ws.cell(r, 4, act); c2.alignment = Alignment(vertical="center", horizontal="right")
    if act == "Card E":
        c2.font = font(10, True, OXBLOOD); c2.hyperlink = "#'E Bank change'!A1"
    else:
        c2.font = font(10, color=MUTED); c2.fill = fill(INPUT_F)
    for col in range(1, 6): ws.cell(r, col).border = bot
    ws.row_dimensions[r].height = 22
    r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
n = ws.cell(r, 1, "Add a row every time you hit a situation the cards don't cover, rather than handling it from memory. "
                  "Rare steps fail because you forget they exist, not because you don't know how to do them.")
n.font = font(9, color=MUTED); n.alignment = WRAP
ws.row_dimensions[r].height = 30
r += 2

r = heading(ws, r, "Standing rules")
RULES = [
 ("Status lives in the system, not in your head",
  "Every invoice sits in exactly one named state at all times. If you get pulled away mid-item, park it in the state "
  "it is actually in before you look at anything else. When you come back you never ask \"where was I\" — you ask the queue."),
 ("Never re-send a payment file from memory",
  "If you can't tell whether a run went out, the bank's confirmation is the only source of truth. Check bank activity "
  "before you touch anything. This is the most expensive interruption failure in the job."),
]
for i, (t, b) in enumerate(RULES):
    col = OXBLOOD if i == 1 else PETROL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(r, 1, t); c.font = font(10.5, True, col); c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r].height = 20; r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(r, 1, b); c.font = font(9.5, color=INK); c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.row_dimensions[r].height = 32; r += 2

ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
c = ws.cell(r, 1, "Once a month, or after any miss: which card should have caught it? If the step was missing, add it. "
                  "If it was there and got skipped, ask whether the card is too long, or whether that step belongs in "
                  "the never-skip list. A checklist you've stopped running is worse than none, because you think you're covered.")
c.font = font(9, color=MUTED, i=True); c.alignment = WRAP
ws.row_dimensions[r].height = 44
ws.freeze_panes = "A3"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ============ Run log ============
lg = wb.create_sheet("Run log")
lg.sheet_view.showGridLines = False
HEADS = ["Run date", "Batch total", "Item count", "Confirmation ref", "Time sent",
         "Released by", "Rejects worked", "Notes"]
WIDTHS = [12, 14, 11, 22, 11, 16, 14, 40]
for i, (h, w) in enumerate(zip(HEADS, WIDTHS), start=1):
    c = lg.cell(1, i, h); c.font = font(9, True, WHITE); c.fill = fill(PETROL)
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    lg.column_dimensions[get_column_letter(i)].width = w
lg.row_dimensions[1].height = 26
EXAMPLE = ["2026-09-04", 48250.00, 37, "BANK-REF-000000", "14:20", "[your name]", "Yes",
           "EXAMPLE ROW — delete it. One row per run; fills in from Card B and Card C."]
for i, v in enumerate(EXAMPLE, start=1):
    c = lg.cell(2, i, v); c.font = font(10, color=MUTED, i=True); c.border = bot
    c.alignment = Alignment(vertical="center", wrap_text=(i == 8))
lg.cell(2, 2).number_format = '$#,##0.00'
lg.cell(2, 3).number_format = '#,##0'
lg.row_dimensions[2].height = 28
for rr in range(3, 41):
    for i in range(1, 9):
        c = lg.cell(rr, i, None); c.border = bot; c.font = font(10)
        c.alignment = Alignment(vertical="center", wrap_text=(i == 8))
    lg.cell(rr, 2).number_format = '$#,##0.00'
    lg.cell(rr, 3).number_format = '#,##0'
    lg.row_dimensions[rr].height = 20
lg.cell(42, 1, "Runs logged:").font = font(9, True, MUTED)
tc = lg.cell(42, 2, "=COUNT($B$3:$B$40)"); tc.font = font(10, True, PETROL)
lg.cell(43, 1, "Total paid:").font = font(9, True, MUTED)
tp = lg.cell(43, 2, "=SUM($B$3:$B$40)"); tp.font = font(10, True, PETROL); tp.number_format = '$#,##0.00'
lg.freeze_panes = "A3"

wb.save(OUT)
print("saved", OUT)
